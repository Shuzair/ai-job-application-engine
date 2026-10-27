#!/usr/bin/env python3
"""The fixed Netherlands layout for the CV and the cover letter.

One layout for every application, not configurable per job: A4, Calibri
11 pt, 2 cm margins, navy headings with a rule, one column with no tables or
text boxes (ATS-safe). The CV's sections and skill categories are fixed for
a data / analytics engineer applying from abroad. Both documents share the
same header, so they match. Only the text changes from job to job; it comes
from cv-data.yaml / cl-data.yaml.

To change the look of every document, change the constants below.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ── The format ──────────────────────────────────────────────────────────────
PAGE_W, PAGE_H = Cm(21.0), Cm(29.7)  # A4
MARGIN = Cm(2.0)
FONT = "Calibri"
BODY_PT = 11
HEADER_PT = 10        # headline, contact, links and relocation lines
NAME_PT = 18
HEADING_PT = 12
LINE = 276            # 1.15 line spacing (240 = single)
NAVY = RGBColor(0x1A, 0x3A, 0x5C)
LINK = "2E6DA4"
SEP = " | "
DOT = " · "

# CV sections in rendered order: (cv-data.yaml key, heading). Skills sit right
# under the summary so a Dutch recruiter sees the stack before the employers,
# which they may not know.
CV_SECTIONS = [
    ("summary", "Professional Summary"),
    ("skills", "Technical Skills"),
    ("experience", "Professional Experience"),
    ("projects", "Projects"),
    ("education", "Education"),
    ("certifications", "Certifications"),
    ("languages", "Languages"),
    ("interests", "Interests"),
]

# Skill categories in rendered order. Mirrored as the `category` enum in
# .claude/skills/cv-style-netherlands/cv-data-schema.yaml.
SKILL_CATEGORIES = [
    "Programming",
    "Data Warehousing & Databases",
    "Pipelines & Orchestration",
    "Data Modelling & BI",
    "Cloud & DevOps",
    "AI & LLMs",
]


# ── Helpers ─────────────────────────────────────────────────────────────────

def _set_spacing(paragraph, before=0, after=0, line=None):
    """Set paragraph spacing in twips."""
    pPr = paragraph._p.get_or_add_pPr()
    for existing in pPr.findall(qn("w:spacing")):
        pPr.remove(existing)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), str(before))
    sp.set(qn("w:after"), str(after))
    if line is not None:
        sp.set(qn("w:line"), str(line))
        sp.set(qn("w:lineRule"), "auto")
    pPr.append(sp)


def _add_bottom_border(paragraph):
    """Add a horizontal rule (bottom border) to a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "6")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), "000000")
    pBdr.append(bot)
    pPr.append(pBdr)


def _para(doc, *, before=0, after=0, line=None, center=False, style=None):
    p = doc.add_paragraph(style=style)
    _set_spacing(p, before, after, line)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _run(paragraph, text, *, size=BODY_PT, bold=False, italic=False, color=None):
    r = paragraph.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = color
    return r


def _link(paragraph, url, text, size=BODY_PT):
    """Insert a clickable hyperlink run into *paragraph*."""
    r_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run_el = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rPr.append(rFonts)
    for tag in ("w:sz", "w:szCs"):
        sz = OxmlElement(tag)
        sz.set(qn("w:val"), str(int(size * 2)))
        rPr.append(sz)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), LINK)
    rPr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    run_el.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run_el.append(t)
    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)


def _display(url):
    return url.replace("https://", "").replace("http://", "").rstrip("/")


def _new_document():
    """A4 document with the fixed margins. Returns (doc, usable width)."""
    doc = Document()
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(BODY_PT)
    bullet = doc.styles["List Bullet"].paragraph_format
    bullet.space_before = Pt(0)
    bullet.space_after = Pt(0)
    s = doc.sections[0]
    s.page_width, s.page_height = PAGE_W, PAGE_H
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = MARGIN
    return doc, s.page_width - s.left_margin - s.right_margin


def _header(doc, hdr, headline):
    """Name, headline, contact, links, relocation — the same on CV and letter."""
    _run(_para(doc, center=True), hdr["name"].upper(), size=NAME_PT, bold=True, color=NAVY)
    if headline:
        _run(_para(doc, center=True), headline, size=HEADER_PT)

    contact = [hdr[k] for k in ("location", "phone", "email") if hdr.get(k)]
    if contact:
        _run(_para(doc, center=True), SEP.join(contact), size=HEADER_PT)

    links = [hdr[k] for k in ("linkedin", "github") if hdr.get(k)]
    if links:
        lp = _para(doc, after=60, center=True)
        for i, url in enumerate(links):
            if i:
                _run(lp, SEP, size=HEADER_PT)
            _link(lp, url, _display(url), size=HEADER_PT)

    if hdr.get("relocation"):
        _run(_para(doc, after=60, center=True), hdr["relocation"], size=HEADER_PT, bold=True)


def _heading(doc, title):
    p = _para(doc, before=200, after=40)
    _run(p, title.upper(), size=HEADING_PT, bold=True, color=NAVY)
    _add_bottom_border(p)


def _bullets(doc, items):
    for b in items:
        _run(_para(doc, line=LINE, style="List Bullet"), b.strip())


def _dated_line(doc, usable_w, before):
    p = _para(doc, before=before)
    p.paragraph_format.tab_stops.add_tab_stop(usable_w, WD_TAB_ALIGNMENT.RIGHT)
    return p


def _dates(paragraph, start, end):
    if start or end:
        paragraph.add_run("\t")
        _run(paragraph, f"{start or ''} – {end or ''}")


# ── CV sections ─────────────────────────────────────────────────────────────

def _summary(doc, summary, _w):
    _run(_para(doc, after=60, line=LINE), summary.strip())


def _skills(doc, groups, _w):
    by_category = {g["category"]: g["items"] for g in groups}
    for category in SKILL_CATEGORIES:
        if by_category.get(category):
            p = _para(doc, after=20, line=LINE)
            _run(p, f"{category}:", bold=True)
            _run(p, f" {DOT.join(by_category[category])}")


def _experience(doc, roles, usable_w):
    for i, role in enumerate(roles):
        p = _dated_line(doc, usable_w, before=120 if i else 0)
        _run(p, role["title"], bold=True)
        _run(p, f"{SEP}{role['company']}")
        where = _where(role)
        if where:
            _run(p, f"{SEP}{where}")
        _dates(p, role.get("date_start"), role.get("date_end"))

        if role.get("company_description"):
            _run(_para(doc, after=40), role["company_description"].strip(), italic=True)
        _bullets(doc, role.get("bullets", []))


def _where(role):
    """'Lahore, Pakistan (Remote, Hoboken, US)' from whichever parts are present."""
    loc = role.get("location") or ""
    extra = [x for x in (role.get("work_model"), role.get("hq")) if x]
    if loc and extra:
        return f"{loc} ({', '.join(extra)})"
    return loc or ", ".join(extra)


def _projects(doc, projects, _w):
    for i, proj in enumerate(projects):
        p = _para(doc, before=120 if i else 0)
        _run(p, proj["name"], bold=True)
        if proj.get("url"):
            _run(p, SEP)
            _link(p, proj["url"], _display(proj["url"]))
        if proj.get("description"):
            _run(_para(doc, line=LINE), proj["description"].strip(), italic=True)
        _bullets(doc, proj.get("bullets", []))


def _education(doc, entries, usable_w):
    for i, edu in enumerate(entries):
        p = _dated_line(doc, usable_w, before=60 if i else 0)
        status = f" ({edu['status']})" if edu.get("status") else ""
        _run(p, f"{edu['degree']}{status}", bold=True)
        extra = [x for x in (edu.get("institution"), edu.get("location")) if x]
        if extra:
            _run(p, SEP + SEP.join(extra))
        _dates(p, edu.get("date_start"), edu.get("date_end"))
        if edu.get("gpa"):
            _run(_para(doc), f"GPA: {edu['gpa']}")


def _certifications(doc, certs, usable_w):
    for cert in certs:
        p = _dated_line(doc, usable_w, before=0)
        _set_spacing(p, after=20, line=LINE)
        _run(p, cert["name"], bold=True)
        _run(p, f"{SEP}{cert['issuer']}")
        if cert.get("url"):
            _run(p, SEP)
            _link(p, cert["url"], "Verify")
        if cert.get("date"):
            p.add_run("\t")
            _run(p, cert["date"])


def _languages(doc, langs, _w):
    p = _para(doc)
    for i, lang in enumerate(langs):
        if i:
            _run(p, DOT)
        _run(p, f"{lang['language']}:", bold=True)
        _run(p, f" {lang['level']}")


def _interests(doc, interests, _w):
    _run(_para(doc), DOT.join(interests))


_SECTION_RENDERERS = {
    "summary": _summary,
    "skills": _skills,
    "experience": _experience,
    "projects": _projects,
    "education": _education,
    "certifications": _certifications,
    "languages": _languages,
    "interests": _interests,
}


# ── Documents ───────────────────────────────────────────────────────────────

def render_cv(data, output_path):
    """Build cv.docx from cv-data.yaml content in the fixed layout."""
    doc, usable_w = _new_document()
    _header(doc, data["header"], SEP.join(data["header"].get("titles", [])))
    for key, heading in CV_SECTIONS:
        if data.get(key):
            _heading(doc, heading)
            _SECTION_RENDERERS[key](doc, data[key], usable_w)
    doc.save(str(output_path))
    return str(output_path)


def render_cl(data, output_path):
    """Build cover-letter.docx from cl-data.yaml content in the fixed layout."""
    doc, _ = _new_document()
    _header(doc, data["header"], data["header"].get("title"))
    _add_bottom_border(_para(doc, before=60, after=120))

    _run(_para(doc, after=0), data["date"])
    _para(doc, line=LINE)  # blank line between date and salutation
    _run(_para(doc, after=120), data["salutation"])
    for text in data["paragraphs"]:
        _run(_para(doc, after=120, line=LINE), text.strip())

    so = data["sign_off"]
    _run(_para(doc, before=120, after=240), so["closing"])
    _run(_para(doc), so["name"], bold=True, color=NAVY)
    for field in ("title", "email", "phone"):
        if so.get(field):
            _run(_para(doc), so[field])

    doc.save(str(output_path))
    return str(output_path)
